import os

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE

from tools.billing import get_bill
from tools.memory import get_preference
from tools.analytics import get_sales_summary
from tools.errors import ToolError

OUTPUT_DIR = os.path.join(os.path.dirname(__file__), "..", "data", "generated")
os.makedirs(OUTPUT_DIR, exist_ok=True)


def generate_invoice_pdf(bill_id):
    data = get_bill(bill_id)
    bill, items = data["bill"], data["items"]
    if bill["status"] != "finalized":
        raise ToolError(f"Bill #{bill_id} isn't finalized yet — finalize the sale before generating an invoice.")

    shop_name = get_preference("shop_name") or "Nebula Kirana Store"
    gstin = get_preference("shop_gstin") or "—"
    address = get_preference("shop_address") or ""

    path = os.path.join(OUTPUT_DIR, f"invoice_{bill_id}.pdf")
    doc = SimpleDocTemplate(path, pagesize=A4, topMargin=18 * mm, bottomMargin=18 * mm,
                             leftMargin=16 * mm, rightMargin=16 * mm)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("ShopTitle", parent=styles["Title"], fontSize=16, spaceAfter=2)
    small = ParagraphStyle("Small", parent=styles["Normal"], fontSize=9)

    story = [
        Paragraph(shop_name, title_style),
        Paragraph(address, small) if address else Spacer(1, 0),
        Paragraph(f"GSTIN: {gstin}", small),
        Spacer(1, 8),
        Paragraph(f"<b>TAX INVOICE</b> &nbsp;&nbsp; Invoice #: {bill_id}", styles["Heading3"]),
        Paragraph(f"Date: {bill['finalized_at']}", small),
        Paragraph(f"Customer: {bill['customer_name'] or 'Walk-in'}", small),
        Paragraph(
            f"Payment: {bill['payment_mode'].upper()}" + (f" (ref: {bill['payment_ref']})" if bill["payment_ref"] else ""),
            small,
        ),
        Spacer(1, 10),
    ]

    table_data = [["Item", "HSN", "Qty", "Rate (Rs.)", "Taxable (Rs.)", "CGST (Rs.)", "SGST (Rs.)", "Amount (Rs.)"]]
    for it in items:
        table_data.append([
            it["product_name"], it["hsn_code"] or "-", f"{it['qty']} {it['unit']}",
            f"{it['unit_price']:.2f}", f"{it['taxable_value']:.2f}",
            f"{it['cgst_amt']:.2f}", f"{it['sgst_amt']:.2f}", f"{it['line_total']:.2f}",
        ])

    n_item_rows = len(items)
    table_data += [
        ["", "", "", "", "Subtotal", "", "", f"{bill['subtotal']:.2f}"],
        ["", "", "", "", "CGST", f"{bill['cgst_total']:.2f}", "", ""],
        ["", "", "", "", "SGST", "", f"{bill['sgst_total']:.2f}", ""],
        ["", "", "", "", "Round off", "", "", f"{bill['round_off']:.2f}"],
        ["", "", "", "", "GRAND TOTAL", "", "", f"Rs. {bill['total']:.2f}"],
    ]

    t = Table(table_data, repeatRows=1, colWidths=[95, 35, 45, 48, 58, 42, 42, 55])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4a2545")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, n_item_rows), 0.5, colors.grey),
        ("ALIGN", (2, 0), (-1, -1), "RIGHT"),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("LINEABOVE", (0, -1), (-1, -1), 0.75, colors.black),
    ]))
    story.append(t)
    story.append(Spacer(1, 16))
    story.append(Paragraph("Thank you for shopping with us. Goods once sold are not taken back.", small))
    doc.build(story)

    return {"path": path, "bill_id": bill_id, "total": bill["total"]}


def generate_analysis_deck(start_date, end_date):
    summary = get_sales_summary(start_date, end_date)
    if summary["bill_count"] == 0:
        raise ToolError(f"No finalized bills between {start_date} and {end_date} — nothing to analyze yet.")

    shop_name = get_preference("shop_name") or "Nebula Kirana Store"

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Title slide
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = f"{shop_name} — Sales Analysis"
    slide.placeholders[1].text = f"{start_date} to {end_date}"

    # KPI slide
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Key Numbers"
    tf = slide.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11), Inches(4.5)).text_frame
    kpis = [
        f"Total sales: Rs. {summary['total_sales']:.2f}",
        f"Bills cut: {summary['bill_count']}",
        f"CGST collected: Rs. {summary['cgst_collected']:.2f}",
        f"SGST collected: Rs. {summary['sgst_collected']:.2f}",
        f"Products below reorder level: {len(summary['low_stock'])}",
    ]
    for i, k in enumerate(kpis):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = k
        p.font.size = Pt(24)

    # Daily sales chart
    if summary["by_day"]:
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = "Sales by Day"
        chart_data = CategoryChartData()
        days = sorted(summary["by_day"].keys())
        chart_data.categories = days
        chart_data.add_series("Sales (Rs.)", [summary["by_day"][d] for d in days])
        slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.5), Inches(11.3), Inches(5.5), chart_data)

    # Top items chart
    if summary["top_items"]:
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = "Top Items by Revenue"
        chart_data = CategoryChartData()
        chart_data.categories = [t["product_name"] for t in summary["top_items"]]
        chart_data.add_series("Revenue (Rs.)", [round(t["revenue"], 2) for t in summary["top_items"]])
        slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(1), Inches(1.5), Inches(11.3), Inches(5.5), chart_data)

    # Reorder attention slide
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Reorder Attention"
    tf = slide.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11), Inches(5)).text_frame
    if summary["low_stock"]:
        for i, item in enumerate(summary["low_stock"]):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = f"{item['name']}: {item['qty_on_hand']}{item['unit']} left (reorder at {item['reorder_level']})"
            p.font.size = Pt(18)
    else:
        tf.paragraphs[0].text = "All stock levels are healthy."
        tf.paragraphs[0].font.size = Pt(20)

    path = os.path.join(OUTPUT_DIR, f"analysis_{start_date}_to_{end_date}.pptx")
    prs.save(path)

    return {"path": path, "start_date": start_date, "end_date": end_date, "total_sales": summary["total_sales"]}
